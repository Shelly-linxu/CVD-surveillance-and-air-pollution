"""Actual candidate-event descriptives and independent person/event audits.
All identity-linked audit rows stay in private/. No arbitrary across-date collapse.
"""
from pathlib import Path
from collections import Counter, defaultdict
from datetime import date, timedelta
import csv, gzip, hashlib, hmac, json, re, os, html
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from prepare_cases import ROOT, PRIVATE, SOURCES, dt, text

OUT=ROOT/'outputs/air_pollution_cc/case_analysis'
TOP=['脑梗死','急性心肌梗死','心绞痛（急性资格待核）','脑出血','蛛网膜下腔出血']
def save(name,rows,fields=None):
    rows=list(rows)
    if not rows and fields is None: return
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields or list(rows[0])); w.writeheader(); w.writerows(rows)
def age_summary(ee):
    a=np.array([int(e['age']) for e in ee]); q=np.quantile(a,[.25,.5,.75])
    return {'n':len(ee),'age_mean':float(a.mean()),'age_sd':float(a.std(ddof=1)) if len(a)>1 else None,'age_q1':float(q[0]),'age_median':float(q[1]),'age_q3':float(q[2]),'age_min':int(a.min()),'age_max':int(a.max()),'male_n':sum(e['sex']=='男' for e in ee),'female_n':sum(e['sex']=='女' for e in ee),'sex_other_n':sum(e['sex'] not in ('男','女') for e in ee),'age18_64_n':sum(int(e['age'])<65 for e in ee),'age65plus_n':sum(int(e['age'])>=65 for e in ee),'warm_n':sum(e['season']=='暖季' for e in ee),'cold_n':sum(e['season']=='冷季' for e in ee),'terminal_review_n':sum(e['review_status']=='已终审' for e in ee)}
def main():
    OUT.mkdir(parents=True,exist_ok=True)
    with gzip.open(PRIVATE/'candidate_events.csv.gz','rt',encoding='utf-8') as f: events=list(csv.DictReader(f))
    byid={e['event_id']:e for e in events}; extras={}; dateqc=Counter(); treatment=Counter(); coordfields=[]; rawduplicates=defaultdict(list)
    salt=(PRIVATE/'id_salt.bin').read_bytes()
    def token(v): return hmac.new(salt,v.encode(),hashlib.sha256).hexdigest()[:24]
    for yr,fn,sheet in SOURCES:
        p=ROOT/'原始数据'/fn
        expected=pd.read_csv(ROOT/'outputs/air_pollution_cc/source_hashes.csv').set_index('year').loc[yr,'sha256']
        hh=hashlib.sha256()
        with p.open('rb') as f:
            for b in iter(lambda:f.read(1024*1024),b''): hh.update(b)
        assert hh.hexdigest()==expected, 'Source changed since preparation'
        wb=load_workbook(p,read_only=True,data_only=True); it=wb[sheet].iter_rows(values_only=True); header=list(next(it)); pos={v:i for i,v in enumerate(header)}
        coordfields.extend({'year':yr,'field':str(h)} for h in header if any(s in str(h).lower() for s in ['经度','纬度','longitude','latitude','坐标']))
        def get(r,*names):
            for n in names:
                if n in pos: return r[pos[n]]
            return None
        for rn,r in enumerate(it,2):
            eid=f'E{yr}_{rn:07d}'; onset=dt(get(r,'发病日期')); dx=dt(get(r,'诊断日期')); reported=dt(get(r,'报告日期')); card=text(get(r,'报告卡编号'))
            pidraw=text(get(r,'身份证号')).upper(); code=text(get(r,'ICD10')).upper()
            if onset and onset.year==yr and re.fullmatch(r'\d{17}[\dX]|\d{15}',pidraw):
                # Exact ICD prefix, not collapsed I21/I22, for source dedup audit.
                parts=[text(get(r,*ns)) for ns in [('现住地址（省）','现住址省'),('现住地址（市）','现住址市'),('现住地址（区县）','现住址区县'),('现住地址（街道/乡）','现住址街道'),('现住地址（居委会/村）','现住址居委会'),('现住地址（门牌号）','现住址号')]]
                rawduplicates[(token('person:'+pidraw),onset.isoformat(),code[:3])].append({'event_id':eid,'card_hash':token('card:'+card) if card else '', 'address_hash':token('location:'+''.join(parts)),'review':text(get(r,'审核状态','卡片状态')),'icd':code})
            if eid not in byid: continue
            e=byid[eid]
            parts=[text(get(r,*ns)) for ns in [('现住地址（区县）','现住址区县'),('现住地址（街道/乡）','现住址街道'),('现住地址（居委会/村）','现住址居委会'),('现住地址（门牌号）','现住址号')]]
            detail='门牌栏有文字' if parts[-1] and parts[-1] not in ('无','不详','未知','-','0') else '有居委会或村文字' if parts[-2] else '仅街道及以上文字'
            x={'event_id':eid,'district':parts[0],'address_detail':detail,'diagnosis_date':dx.isoformat() if dx else '', 'report_date':reported.isoformat() if reported else '', 'onset_to_diagnosis':(dx-onset).days if dx else '', 'onset_to_report':(reported-onset).days if reported else ''}
            extras[eid]=x
            for kind,val in [('诊断日期',dx),('报告日期',reported)]:
                dateqc[(yr,kind,'缺失或无效' if val is None else '早于发病' if val<onset else '日期逻辑正常')]+=1
            if code.startswith('I20'):
                t=text(get(r,'心绞痛治疗措施'))
                # Do not export free-text clinical content into public summaries.
                cat='缺失' if not t else '提及PCI/介入' if any(k in t.upper() for k in ['PCI','介入','支架']) else '提及CABG/搭桥' if any(k in t.upper() for k in ['CABG','搭桥']) else '有填写（需按原表核实）'
                treatment[(yr,code,cat)]+=1
        wb.close(); print('Source audit finished '+str(yr),flush=True)
    assert len(extras)==len(events)
    enriched=[dict(e,**{k:v for k,v in extras[e['event_id']].items() if k!='event_id'}) for e in events]
    with gzip.open(PRIVATE/'enriched_candidate_events.csv.gz','wt',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(enriched[0])); w.writeheader(); w.writerows(enriched)
    os.chmod(PRIVATE/'enriched_candidate_events.csv.gz',0o600)
    table1=[]
    for d in TOP:
        for yr in ['合计','2023','2024','2025']:
            ee=[e for e in enriched if e['disease']==d and (yr=='合计' or e['year']==yr)]
            if ee: table1.append({'disease':d,'year':yr,**age_summary(ee),'population':'候选事件，定位及急性资格待核'})
    save('表1_候选事件特征.csv',table1)
    sensitivity=[]
    for d in TOP:
        ee=[e for e in enriched if e['disease']==d]
        first={}
        for e in sorted(ee,key=lambda e:(e['onset'],e['event_id'])): first.setdefault(e['person_id'],e)
        sensitivity.append({'disease':d,'all_candidate_n':len(ee),'terminal_review_only_n':sum(e['review_status']=='已终审' for e in ee),'first_record_per_person_disease_n':len(first),'loss_if_first_only_n':len(ee)-len(first),'I20_0_n':sum(e['icd10'].startswith('I20.0') for e in ee) if d==TOP[2] else None})
    save('纳入口径敏感性计数.csv',sensitivity)
    save('发病诊断报告日期质量.csv',[{'year':yr,'field':k,'status':s,'n':n} for (yr,k,s),n in sorted(dateqc.items())])
    save('心绞痛编码和治疗字段核查.csv',[{'year':yr,'icd10':c,'treatment_field':t,'n':n} for (yr,c,t),n in sorted(treatment.items())])
    addr=Counter((e['year'],e['address_detail']) for e in enriched)
    save('地址文字完整度.csv',[{'year':yr,'text_detail':s,'n':n,'note':'文字完整度不能证明定位精度'} for (yr,s),n in sorted(addr.items())])
    groups=defaultdict(list)
    for e in enriched: groups[(e['person_id'],e['icd10'][:3])].append(e)
    repeat=[]; intervals=Counter()
    for (pid,code),ee in groups.items():
        ee.sort(key=lambda e:(e['onset'],e['event_id']))
        for prev,cur in zip(ee,ee[1:]):
            gap=(date.fromisoformat(cur['onset'])-date.fromisoformat(prev['onset'])).days
            band='同日' if gap==0 else '1–28天（仅核查，未自动合并）' if gap<=28 else '29–90天' if gap<=90 else '91–365天' if gap<=365 else '>365天'
            intervals[(code,band)]+=1
            if gap<=28: repeat.append({'person_id':pid,'icd_prefix':code,'previous_event_id':prev['event_id'],'current_event_id':cur['event_id'],'days':gap,'status':'待核查，未判定重复或复发'})
    save('同人同病种事件间隔.csv',[{'icd_prefix':c,'interval':b,'adjacent_pairs':n} for (c,b),n in sorted(intervals.items())])
    fields=['person_id','icd_prefix','previous_event_id','current_event_id','days','status']
    with (PRIVATE/'short_interval_event_audit.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(repeat)
    os.chmod(PRIVATE/'short_interval_event_audit.csv',0o600)
    duplicates=[v for v in rawduplicates.values() if len(v)>1]
    dupstats={'same_person_same_date_same_icd_groups':len(duplicates),'rows_in_these_groups':sum(len(v) for v in duplicates),'groups_with_different_address':sum(len(set(x['address_hash'] for x in v))>1 for v in duplicates),'groups_with_different_full_icd':sum(len(set(x['icd'] for x in v))>1 for v in duplicates),'short_interval_pairs_in_candidates':len(repeat),'coordinate_columns_found':len(coordfields)}
    (OUT/'身份证事件重复核查汇总.json').write_text(json.dumps(dupstats,ensure_ascii=False,indent=2))
    daily=Counter((e['onset'],e['disease']) for e in enriched if e['disease'] in TOP)
    date_range=pd.date_range('2023-01-01','2025-12-31')
    dailyrows=[{'date':d.date().isoformat(),'disease':dis,'candidate_count':daily[(d.date().isoformat(),dis)]} for d in date_range for dis in TOP]
    save('五类疾病逐日候选事件数.csv',dailyrows)
    frame=pd.DataFrame(dailyrows); frame['month']=frame['date'].str[:7]
    monthly=frame.groupby(['month','disease'],sort=True)['candidate_count'].sum().reset_index(); monthly.to_csv(OUT/'五类疾病逐月候选事件数.csv',index=False,encoding='utf-8-sig')
    import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
    from matplotlib.font_manager import FontProperties
    font=FontProperties(fname='/System/Library/Fonts/STHeiti Medium.ttc')
    colors=['#126782','#e67e22','#7b5ea7','#2a9d8f','#a63a50']
    fig,axes=plt.subplots(5,1,figsize=(12,11),sharex=True)
    for ax,d,c in zip(axes,TOP,colors):
        sub=monthly[monthly.disease==d]; ax.plot(pd.to_datetime(sub.month),sub.candidate_count,color=c,marker='o',markersize=3)
        ax.set_ylabel('候选事件数',fontproperties=font); ax.set_title(d,fontproperties=font,loc='left',fontsize=11); ax.grid(axis='y',alpha=.2)
    fig.suptitle('2023—2025年五类疾病逐月候选事件数',fontproperties=font,fontsize=16)
    fig.text(.08,.012,'计数未经地址定位和急性资格最终审定；年度变化不能直接解释为发病率变化。',fontproperties=font,fontsize=10)
    fig.tight_layout(rect=[0,.035,1,.97]); fig.savefig(OUT/'五类疾病逐月候选事件数.png',dpi=180); plt.close(fig)
    # Report only aggregate findings; identifier-linked audit files remain private.
    def t(rows,cols):
        return '<table><tr>'+''.join('<th>'+html.escape(c)+'</th>' for c in cols)+'</tr>'+''.join('<tr>'+''.join('<td>'+html.escape(str(r[c]))+'</td>' for c in cols)+'</tr>' for r in rows)+'</table>'
    base=[{'疾病':r['disease'],'候选事件数':r['n'],'年龄均值':round(r['age_mean'],1),'年龄中位数':r['age_median'],'男性':r['male_n'],'女性':r['female_n'],'≥65岁':r['age65plus_n'],'已终审':r['terminal_review_n']} for r in table1 if r['year']=='合计']
    content=f'''<h1>病例描述统计与身份证事件核查</h1><p>2026年10月2日。已直接复核三个年度Excel，源文件SHA256与前期整理一致。统计单位是候选事件，未完成地理编码和正式关联模型。</p><h2>五类疾病候选事件特征</h2>{t(base,list(base[0]))}<p>年龄按出生日期至发病日计算，未加入病例交叉模型中的年龄和性别固定主效应。心绞痛的ICD细分和治疗字段已另表核查；治疗字段有填写不能自动证明急性资格。</p><h2>时间分布</h2><img src="五类疾病逐月候选事件数.png" style="width:100%"><p>零事件日期已补齐。2023年监测系统在4月投入使用、11月联通随访，来自原始2023年监测报告。其影响需检查报告日期及补发来源，不能根据当前曲线事后选择排除月份。</p><h2>身份证重复事件核查</h2><p>原始资料中，同身份证、同发病日、同ICD前三位存在{dupstats['same_person_same_date_same_icd_groups']}个重复候选组，共{dupstats['rows_in_these_groups']}条记录，其中{dupstats['groups_with_different_address']}组地址不同、{dupstats['groups_with_different_full_icd']}组完整ICD不同。此前候选表已初步合并同卡及同人同日同类，最终事件需复核这些冲突，不仅凭身份证删除所有重复出现的个人。</p><p>候选事件中发现{len(repeat)}对同身份证、同ICD前三位且相邻发病日期间隔≤28天的记录。这些仅标为待核查，未自动合并；不同日期可能是真实复发。已计算每人每病种监测期首条记录的敏感性样本数。2023原始官方报告明确使用28天监测重卡规则，但不能据此假定三个年度所有疾病的研究事件均应自动按28天合并。</p><h2>地址和污染读取</h2><p>三个病例表未发现经纬度字段。地址文字完整度已统计，实际坐标及精度仍待获得。CHAP样本文件核实为uint16、缺失值65535、比例因子0.1，实际读取须自动处理比例因子和缺失值，不能将原始整数当作μg/m³浓度。</p><h2>当前分析边界</h2><p>本报告为实际病例描述统计与质量核查。尚无个体污染匹配、逐日气象和正式纳入审定，未生成OR、置信区间或P/q值。高德个人实名认证已完成，研究应用和Web服务Key已创建；测试住址发送授权仍待确认，尚未执行地址请求。未向地图平台发送病例身份或临床信息。最终需使用可验证定位至约1km的住址，低精度地址另作敏感性分析。</p>'''
    (OUT/'病例分析报告.html').write_text('<!doctype html><html lang="zh"><meta charset="utf-8"><title>病例分析报告</title><style>body{max-width:1050px;margin:36px auto;padding:20px;font:16px/1.7 system-ui;color:#203448}table{border-collapse:collapse;width:100%;font-size:14px}td,th{padding:9px;border-bottom:1px solid #dbe3e8;text-align:left}th{background:#eef4f7}h2{color:#126782;margin-top:30px}</style><body>'+content+'</body></html>',encoding='utf-8')
    (OUT/'analysis_summary.json').write_text(json.dumps({'candidate_events':len(events),'top5_candidate_events':sum(r['n'] for r in table1 if r['year']=='合计'),'daily_series_dates':len(date_range),'duplicate_audit':dupstats,'association_models_fitted':0},ensure_ascii=False,indent=2))
    print(json.dumps(dupstats,ensure_ascii=False),flush=True)
if __name__=='__main__': main()
