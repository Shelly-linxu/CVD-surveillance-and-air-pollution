"""Aggregate ID quality audit. Never export raw national IDs or names."""
from collections import Counter
from datetime import date
from pathlib import Path
import csv,json,re
from openpyxl import load_workbook
from prepare_cases import ROOT,SOURCES,text,dt
OUT=ROOT/'outputs/air_pollution_cc/case_analysis'
WEIGHTS=[7,9,10,5,8,4,2,1,6,3,7,9,10,5,8,4,2]
CHECK='10X98765432'

def classify(value):
    s=text(value).upper()
    if re.fullmatch(r'\d{15}',s):
        expanded=s[:6]+'19'+s[6:]
        expanded+=CHECK[sum(int(x)*w for x,w in zip(expanded,WEIGHTS))%11]
        label='15位旧证件格式（可转换）'
    elif re.fullmatch(r'\d{17}[\dX]',s):
        expanded=s
        if CHECK[sum(int(x)*w for x,w in zip(s[:17],WEIGHTS))%11]!=s[-1]: return '18位格式但校验码不符',None,None
        label='18位且校验码符合'
    else: return '缺失或格式不符',None,None
    try: dob=date(int(expanded[6:10]),int(expanded[10:12]),int(expanded[12:14]))
    except ValueError:return '证件内出生日期无效',None,None
    return label,dob,'男' if int(expanded[16])%2 else '女'

def main():
    counts=Counter()
    for year,name,sheet in SOURCES:
        wb=load_workbook(ROOT/'原始数据'/name,read_only=True,data_only=True)
        it=wb[sheet].iter_rows(values_only=True); header=list(next(it));pos={v:i for i,v in enumerate(header)}
        for row in it:
            if not any(x is not None for x in row): continue
            value=row[pos['身份证号']]
            label,dob,sex=classify(value)
            counts[(year,'身份证格式和校验',label)]+=1
            if isinstance(value,(int,float)) and not isinstance(value,bool): counts[(year,'存储类型','数值型身份证（存在精度风险，需原证核实）')]+=1
            if dob:
                recorded=dt(row[pos['出生日期']])
                if recorded: counts[(year,'身份证与出生日期','一致' if dob==recorded else '不一致')]+=1
                else:counts[(year,'身份证与出生日期','出生日期缺失')]+=1
                recordedsex=text(row[pos['性别']])
                counts[(year,'身份证与性别','一致' if recordedsex==sex else '不一致或未标明')]+=1
        wb.close();print('Identity audit completed:',year,flush=True)
    rows=[dict(year=y,audit=a,result=r,n=n) for (y,a,r),n in sorted(counts.items())]
    with (OUT/'身份证格式校验与人口学一致性.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=['year','audit','result','n']);w.writeheader();w.writerows(rows)
    (OUT/'identity_audit_metadata.json').write_text(json.dumps(dict(scope='全部原始非空记录，未按纳入规则过滤',effect='仅核查标记，未自动排除校验不符的记录；不代表证件真实性验证',note='不输出姓名或原始身份证号码；15位证件按旧证件规则转换仅用于格式核查'),ensure_ascii=False,indent=2))

if __name__=='__main__':main()
