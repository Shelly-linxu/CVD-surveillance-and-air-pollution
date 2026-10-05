"""Read source workbooks without alteration; create LOCAL candidate data, not final eligibility."""
from pathlib import Path
from datetime import datetime, date
import calendar, csv, gzip, hashlib, hmac, json, os, re
from collections import Counter, defaultdict
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs/air_pollution_cc'
PRIVATE = ROOT / 'work/air_pollution_cc/private'
# Configure source workbook names locally; never commit this configuration.
_config = ROOT / 'config/sources.local.json'
SOURCES = [(int(s['year']), s['filename'], s['sheet']) for s in json.loads(_config.read_text())['sources']] if _config.exists() else []

MAP = {'I20':'心绞痛（急性资格待核）', 'I21':'急性心肌梗死', 'I22':'急性心肌梗死', 'I46':'心脏性猝死（定义待核）', 'I60':'蛛网膜下腔出血', 'I61':'脑出血', 'I63':'脑梗死', 'I64':'未特指卒中'}

def text(v): return '' if v is None else re.sub(r'\s+', '', str(v))
def dt(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    # 2023 dates are unformatted numeric Excel serials (1900 date system).
    if isinstance(v,(int,float)) and 1 <= v <= 80000:
        try: return from_excel(v).date()
        except (ValueError,OverflowError,AttributeError): return None
    s=text(v)
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y%m%d', '%Y-%m-%d%H:%M:%S', '%Y/%m/%d%H:%M:%S'):
        try: return datetime.strptime(s, fmt).date()
        except ValueError: pass
    return None
def write(name, rows, fields=None):
    rows=list(rows)
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f, fieldnames=fields or list(rows[0])); w.writeheader(); w.writerows(rows)
def candidates(d):
    return [date(d.year,d.month,n) for n in range(1,calendar.monthrange(d.year,d.month)[1]+1) if date(d.year,d.month,n).weekday()==d.weekday()]

def main():
    if not SOURCES: raise RuntimeError('Create config/sources.local.json from the example before running')
    OUT.mkdir(parents=True,exist_ok=True); PRIVATE.mkdir(parents=True,exist_ok=True); os.chmod(PRIVATE,0o700)
    saltfile=PRIVATE/'id_salt.bin'
    if not saltfile.exists(): saltfile.write_bytes(os.urandom(32)); os.chmod(saltfile,0o600)
    salt=saltfile.read_bytes()
    def token(s): return hmac.new(salt,s.encode(),hashlib.sha256).hexdigest()[:24]
    stats=[]; fieldrows=[]; catrows=[]; events=[]; addresses={}; seen_card=set(); seen_event=set(); audit=[]; hashes=[]
    for yr,fn,sheet in SOURCES:
        path=ROOT/'原始数据'/fn
        hh=hashlib.sha256()
        with path.open('rb') as f:
            for chunk in iter(lambda:f.read(1024*1024),b''): hh.update(chunk)
        hashes.append({'year':yr,'filename':fn,'sheet':sheet,'sha256':hh.hexdigest()})
        wb=load_workbook(path,read_only=True,data_only=True); ws=wb[sheet]; it=ws.iter_rows(values_only=True); headers=list(next(it)); idx={v:i for i,v in enumerate(headers)}
        for i,h in enumerate(headers): fieldrows.append({'year':yr,'column':i+1,'field':h})
        counts=Counter(); cats=defaultdict(Counter)
        def get(row,*names):
            for n in names:
                if n in idx: return row[idx[n]]
            return None
        for rowno,row in enumerate(it,2):
            if all(v is None for v in row): continue
            counts['原始非空记录']+=1
            for key,names in {'重卡标记':('是否重卡',),'审核状态':('审核状态','卡片状态'),'ICD前三位':('ICD10',),'现住市':('现住地址（市）','现住址市'),'首次标记':('是否首次发病',)}.items():
                val=text(get(row,*names)); cats[key][val[:3] if key=='ICD前三位' else val]+=1
            onset=dt(get(row,'发病日期')); birth=dt(get(row,'出生日期'))
            rawcode=text(get(row,'ICD10')).upper(); m=re.match(r'(I\d{2})',rawcode); code=m.group(1) if m else ''
            age=onset.year-birth.year-((onset.month,onset.day)<(birth.month,birth.day)) if onset and birth else None
            city=text(get(row,'现住地址（市）','现住址市'))
            addrparts=[text(get(row,*ns)) for ns in [('现住地址（省）','现住址省'),('现住地址（市）','现住址市'),('现住地址（区县）','现住址区县'),('现住地址（街道/乡）','现住址街道'),('现住地址（居委会/村）','现住址居委会'),('现住地址（门牌号）','现住址号')]]
            # Do not substitute household registration for current residence.
            pidraw=text(get(row,'身份证号')).upper(); validid=bool(re.fullmatch(r'\d{17}[\dX]|\d{15}',pidraw))
            if not onset or onset.year!=yr: reason='缺失无效或非对应年度发病日期'
            elif age is None or not 0<=age<=120: reason='出生日期缺失或年龄异常'
            elif age<18: reason='年龄小于18岁'
            elif city not in ('广州','广州市'): reason='现住市非广州或缺失'
            elif code not in MAP: reason='未映射ICD类别'
            elif not validid: reason='缺少格式有效的身份证标识（待其他证件链接）'
            elif not any(addrparts[3:]): reason='缺少街道及以下现住址'
            else: reason=''
            if reason: counts[reason]+=1; continue
            pid=token('person:'+pidraw); card=text(get(row,'报告卡编号')); eventkey=(pid,onset.isoformat(),MAP[code])
            cardkey=(yr,card) if card else None
            if cardkey and cardkey in seen_card: counts['重复报告卡号候选（待复核）']+=1; continue
            if eventkey in seen_event: counts['同人同日同类重复候选']+=1; continue
            if cardkey: seen_card.add(cardkey)
            seen_event.add(eventkey)
            loc=token('location:'+''.join(addrparts))
            addresses[loc]={'location_id':loc,'address':''.join(addrparts),'has_door_text':int(bool(addrparts[-1])),'longitude':'','latitude':'','crs':'','precision':'','verified_within_1km':''}
            flags=['event_dedup_pending','geocode_pending','acute_definition_pending','review_status_pending']
            if code=='I20': flags.append('I20_acute_eligibility_pending')
            e={'event_id':f'E{yr}_{rowno:07d}','person_id':pid,'location_id':loc,'onset':onset.isoformat(),'year':yr,'disease':MAP[code],'icd10':rawcode,'age':age,'sex':text(get(row,'性别')),'season':'暖季' if 4<=onset.month<=9 else '冷季','duplicate_flag':text(get(row,'是否重卡')),'review_status':text(get(row,'审核状态','卡片状态')),'first_onset_flag':text(get(row,'是否首次发病')),'has_door_text':int(bool(addrparts[-1])),'status':'candidate_only','pending':';'.join(flags)}
            events.append(e); counts['保留候选记录（非最终事件）']+=1
        wb.close()
        for key,n in counts.items(): stats.append({'year':yr,'step':key,'n':n})
        for key,c in cats.items():
            for value,n in c.items(): catrows.append({'year':yr,'field':key,'value':value or '(空)','n':n})
        print(json.dumps({'year':yr,'flow':dict(counts)},ensure_ascii=False),flush=True)
    write('source_hashes.csv',hashes); write('field_inventory.csv',fieldrows); write('candidate_flow.csv',stats); write('categorical_audit.csv',catrows)
    ef=list(events[0])
    with gzip.open(PRIVATE/'candidate_events.csv.gz','wt',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=ef); w.writeheader(); w.writerows(events)
    with (PRIVATE/'geocoding_queue.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(next(iter(addresses.values())))); w.writeheader(); w.writerows(addresses.values())
    counts=Counter(); nmatch=0; person=Counter(e['person_id'] for e in events)
    with gzip.open(PRIVATE/'candidate_matches.csv.gz','wt',encoding='utf-8',newline='') as f:
        fields=['event_id','person_id','location_id','date','case','disease','age','sex','season','status']; w=csv.DictWriter(f,fieldnames=fields); w.writeheader()
        for e in events:
            d=date.fromisoformat(e['onset']); ds=candidates(d)
            assert len(ds) in (4,5) and sum(x==d for x in ds)==1
            for cd in ds:
                w.writerow({**{k:e[k] for k in ['event_id','person_id','location_id','disease','age','sex','season','status']},'date':cd.isoformat(),'case':int(cd==d)})
            nmatch+=len(ds); counts[len(ds)-1]+=1
    summaries=[]
    for dis in MAP.values():
        if any(s['disease']==dis for s in summaries): continue
        ee=[e for e in events if e['disease']==dis]
        summaries.append({'disease':dis,'candidate_n':len(ee),'n2023':sum(e['year']==2023 for e in ee),'n2024':sum(e['year']==2024 for e in ee),'n2025':sum(e['year']==2025 for e in ee),'male_n':sum(e['sex'] in ('男','1') for e in ee),'female_n':sum(e['sex'] in ('女','2') for e in ee),'age65plus_n':sum(e['age']>=65 for e in ee),'door_text_n':sum(e['has_door_text'] for e in ee),'status':'候选记录，尚非最终表1'})
    summaries.sort(key=lambda x:-x['candidate_n']); write('candidate_disease_summary.csv',summaries)
    summary={'candidate_events':len(events),'candidate_persons':len(person),'persons_with_multiple_candidates':sum(n>1 for n in person.values()),'locations_to_geocode':len(addresses),'candidate_match_rows':nmatch,'groups_by_control_count':dict(counts),'final_eligible_events':None,'models_fitted':0}
    (OUT/'preparation_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    for p in PRIVATE.iterdir(): os.chmod(p,0o600)
    print(json.dumps(summary,ensure_ascii=False),flush=True)

if __name__=='__main__': main()
