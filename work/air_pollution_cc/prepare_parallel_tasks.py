"""Tasks 1/3/4 only: outcome coding review, temporal audit, prospective core plan.
No person-based duplicate checks or event consolidation are performed here.
"""
from pathlib import Path
from datetime import datetime
from collections import Counter
import csv,gzip,json,os,hashlib,html,re,sys
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from prepare_cases import SOURCES,ROOT,dt,text
P=ROOT/'work/air_pollution_cc/private';O=ROOT/'outputs/air_pollution_cc/parallel_preparation';O.mkdir(parents=True,exist_ok=True)
OUTCOMES=[('I63','脑梗死',['I63']),('AMI','急性心肌梗死',['I21','I22']),('I20_0','不稳定型心绞痛',['I20.0']),('I61','脑出血',['I61']),('I60','蛛网膜下腔出血',['I60'])]
def outcome(code):return next((key for key,name,prefs in OUTCOMES if any(code.startswith(p) for p in prefs)),'')
def save(df,name,private=False):
    path=(P if private else O)/name;df.to_csv(path,index=False,encoding='utf-8-sig');
    if private:os.chmod(path,0o600)
def main():
    events=pd.read_csv(P/'enriched_candidate_events.csv.gz',dtype=str,keep_default_na=False);ids=set(events.event_id)
    hashes=pd.read_csv(ROOT/'outputs/air_pollution_cc/source_hashes.csv').set_index('year');raw=[];clinical={};counts=Counter()
    if '--from-cache' in sys.argv:
        raw=pd.read_csv(P/'raw_temporal_audit.csv.gz',dtype=str,keep_default_na=False).to_dict('records')
        cc=pd.read_csv(P/'event_outcome_candidates.csv.gz',dtype=str,keep_default_na=False).to_dict('records');clinical={r['event_id']:r for r in cc}
    for year,fn,sheet in ([] if '--from-cache' in sys.argv else SOURCES):
        file=ROOT/'原始数据'/fn;assert hashlib.sha256(file.read_bytes()).hexdigest()==hashes.loc[year,'sha256']
        wb=load_workbook(file,read_only=True,data_only=True);it=wb[sheet].iter_rows(values_only=True);head=list(next(it));idx={x:i for i,x in enumerate(head)}
        def get(row,key):return row[idx[key]] if key in idx else None
        for rn,row in enumerate(it,2):
            if all(v is None for v in row):continue
            eid=f'E{year}_{rn:07d}';onset=dt(get(row,'发病日期'));reported=dt(get(row,'报告日期'));diagnosed=dt(get(row,'诊断日期'));code=text(get(row,'ICD10')).upper();diagnosis=text(get(row,'诊断'))
            raw.append(dict(event_id=eid,source_year=year,onset=onset.isoformat() if onset else '',report_date=reported.isoformat() if reported else '',diagnosis_date=diagnosed.isoformat() if diagnosed else '',icd10=code,outcome_id=outcome(code),candidate=eid in ids))
            if eid in ids:
                key=outcome(code);flag='code_supported_text_not_contradictory'
                if key=='I20_0':flag='generic_angina_text_requires_subtype_review' if '不稳定' not in diagnosis else 'unstable_text_agrees'
                if key=='I20_0' and re.search(r'(?<!不)稳定',diagnosis):flag='code_text_conflict'
                if key=='I63' and ('脑出血' in diagnosis or '蛛网膜' in diagnosis):flag='code_text_conflict'
                if key=='I61' and ('脑梗死' in diagnosis or '蛛网膜' in diagnosis):flag='code_text_conflict'
                if not key:flag='outside_five_prespecified_code_groups'
                clinical[eid]=dict(event_id=eid,proposed_outcome=key,icd10=code,diagnosis_text=diagnosis,clinical_review_flag=flag,final_eligibility='pending_final_event_and_address_adjudication')
        wb.close();print('Raw outcome/time audit completed:',year,flush=True)
    raw=pd.DataFrame(raw);assert len(raw)==194925 and len(clinical)==len(events)
    save(raw,'raw_temporal_audit.csv.gz',True);annotation=pd.DataFrame(clinical.values());save(annotation,'event_outcome_candidates.csv.gz',True)
    enriched=events.merge(annotation[['event_id','proposed_outcome','clinical_review_flag']],on='event_id',validate='one_to_one')
    summary=enriched.groupby(['year','proposed_outcome','clinical_review_flag']).size().rename('candidate_events').reset_index();save(summary,'结局编码与文字一致性.csv')
    labels={k:n for k,n,p in OUTCOMES};table=enriched[enriched.proposed_outcome!=''].groupby(['proposed_outcome','year']).size().unstack(fill_value=0).reindex(labels).fillna(0).astype(int);table['total']=table.sum(axis=1);table.insert(0,'outcome_name',[labels[k] for k in table.index]);save(table.reset_index(),'五类预设结局候选计数.csv')
    mapping=pd.DataFrame([dict(outcome_id=k,outcome_name=n,icd_prefixes='|'.join(p),classification='ICD编码候选，非最终合格事件',acute_requirement='真实急性发生日期；编码和诊断文字冲突需复核；不同事件的去重留待定位完成',source='https://klassifikationen.bfarm.de/icd-10-who/kode-suche/htmlamtl2019/block-i20-i25.htm' if k in ('AMI','I20_0') else 'https://klassifikationen.bfarm.de/icd-10-who/kode-suche/htmlamtl2019/block-i60-i69.htm') for k,n,p in OUTCOMES]);save(mapping,'预设结局映射.csv')
    for frame in (raw,enriched):
        for v in ['onset','report_date','diagnosis_date']:frame[v]=pd.to_datetime(frame[v],errors='coerce')
        frame['report_delay_days']=(frame.report_date-frame.onset).dt.days;frame['diagnosis_delay_days']=(frame.diagnosis_date-frame.onset).dt.days
    ds=[]
    for pop,frame,yearcol in [('原始记录未去重',raw,'source_year'),('既有候选事件（未新增去重）',enriched,'year')]:
        for year,sub in frame.groupby(yearcol):
            delay=sub.report_delay_days.dropna();dx=sub.diagnosis_delay_days.dropna()
            ds.append(dict(population=pop,year=year,n=len(sub),onset_missing=int(sub.onset.isna().sum()),report_missing=int(sub.report_date.isna().sum()),report_before_onset=int((delay<0).sum()),diagnosis_before_onset=int((dx<0).sum()),report_delay_median=float(delay.median()),report_delay_p95=float(delay.quantile(.95)),report_delay_max=float(delay.max()),report_delay_over30=int((delay>30).sum()),report_delay_over90=int((delay>90).sum()),onset_diagnosis_report_identical=int(((sub.onset==sub.diagnosis_date)&(sub.onset==sub.report_date)).sum())))
    save(pd.DataFrame(ds),'发病诊断报告时序汇总.csv')
    full_dates=pd.date_range('2023-01-01','2025-12-31');daily=[];flags=[]
    for pop,frame,keycol in [('原始记录未去重',raw,'outcome_id'),('既有候选事件（未新增去重）',enriched,'proposed_outcome')]:
        for key,name,prefs in OUTCOMES:
            sub=frame[frame[keycol]==key];series=sub.groupby('onset').size().reindex(full_dates,fill_value=0)
            for d,n in series.items():daily.append(dict(population=pop,outcome_id=key,outcome_name=name,date=d.strftime('%Y-%m-%d'),n=int(n)))
            for year in [2023,2024,2025]:
                z=series[series.index.year==year];med=float(z.median());mad=float(np.median(abs(z-med)));threshold=max(20,3*med,med+6*1.4826*mad)
                for d,n in z[z>threshold].items():flags.append(dict(population=pop,outcome_id=key,date=d.strftime('%Y-%m-%d'),count=int(n),flag='发病日期计数超过年度稳健筛查阈值',threshold=round(threshold,2),action='仅标记，不能据此直接排除病例或日期'))
        late=frame[(frame.report_delay_days>30)&frame.report_date.notna()]
        grouped=late.groupby('report_date').size()
        for d,n in grouped[grouped>=20].items():flags.append(dict(population=pop,outcome_id='ALL',date=d.strftime('%Y-%m-%d'),count=int(n),flag='同一报告日集中记录至少20条发病超过30天的事件',threshold=20,action='疑似集中补报，需结合系统记录复核；不自动排除'))
    daily=pd.DataFrame(daily);save(daily,'预设五类结局逐日计数_原始与候选.csv');daily['month']=daily.date.str[:7];monthly=daily.groupby(['population','outcome_id','outcome_name','month'],as_index=False).n.sum();save(monthly,'预设五类结局逐月计数.csv')
    flagdf=pd.DataFrame(flags,columns=['population','outcome_id','date','count','flag','threshold','action']);save(flagdf,'延迟报告描述性筛查日期.csv')
    flags=[r for r in flags if r['flag']=='发病日期计数超过年度稳健筛查阈值']
    reportdays=[]
    for pop,frame in [('原始记录未去重',raw),('既有候选事件（未新增去重）',enriched)]:
        for d,sub in frame[frame.report_date.notna()].groupby('report_date'):
            reportdays.append(dict(population=pop,report_date=d.strftime('%Y-%m-%d'),n=len(sub),late_over30=int((sub.report_delay_days>30).sum()),late_over90=int((sub.report_delay_days>90).sum()),delay_median=float(sub.report_delay_days.median()),delay_p95=float(sub.report_delay_days.quantile(.95))))
    reportdf=pd.DataFrame(reportdays);save(reportdf,'逐报告日计数与补报延迟.csv')
    for pop,sub in reportdf.groupby('population'):
        sub=sub.copy();sub['rd']=pd.to_datetime(sub.report_date);sub=sub[sub.rd<=pd.Timestamp(datetime.now().date())]
        for year,y in sub.groupby(sub.rd.dt.year):
            end=min(pd.Timestamp(f'{year}-12-31'),y.rd.max());days=pd.date_range(f'{year}-01-01',end);z=y.set_index('rd').n.reindex(days,fill_value=0);med=float(z.median());mad=float(np.median(abs(z-med)));threshold=max(100,med+6*1.4826*mad)
            peaks=y[(y.n>threshold)&(y.late_over30>=100)&(y.late_over30/y.n>=.75)]
            for row in peaks.to_dict('records'):flags.append(dict(population=pop,outcome_id='ALL',date=row['report_date'],count=row['n'],flag='报告计数超过年度稳健阈值且至少75%延迟超过30天',threshold=round(threshold,2),action='集中补报核查候选；需核对系统日志，不自动排除'))
    flagdf=pd.DataFrame(flags,columns=['population','outcome_id','date','count','flag','threshold','action']);save(flagdf,'监测异常筛查日期清单.csv')
    future=[]
    cutoffs={2023:pd.Timestamp(datetime.now().date()),2024:pd.Timestamp('2026-03-23'),2025:pd.Timestamp('2026-08-17')}
    for pop,frame,ycol in [('原始记录未去重',raw,'source_year'),('既有候选事件（未新增去重）',enriched,'year')]:
        for yr,y in frame.groupby(ycol):
            for field in ['onset','report_date','diagnosis_date']:future.append(dict(population=pop,source_year=yr,field=field,after_extract_cutoff=int((y[field]>cutoffs[int(yr)]).sum()),cutoff=str(cutoffs[int(yr)].date()),cutoff_basis='源文件名下载日期' if int(yr)!=2023 else '源文件无明确下载日期，采用本次核查日期'))
    save(pd.DataFrame(future),'日期超出提取时点核查.csv')
    protocol=dict(version='preanalysis_core_v1',prepared_at=datetime.now().astimezone().isoformat(),association_results_inspected=False,core_parameters_status='fixed_before_real_association_estimation',final_event_eligibility_status='pending; identity dedup expressly deferred until geocoding finishes',design='time-stratified case-crossover; same year/month/day-of-week',years=[2023,2024,2025],age_min=18,outcomes=[dict(outcome_id=k,outcome_name=n,icd_prefixes=p) for k,n,p in OUTCOMES],primary_pollutants=['PM25','PM10','O3'],ozone_metric='MDA8',primary_window='lag01',secondary_windows=['lag0','lag1','lag2','lag3','lag02','lag03'],sensitivity_windows=['lag1','lag12'],pollutant_increment_ug_m3=10,temperature=dict(main_max_lag=21,variable_df=3,lag_df=3,variable_intercept=False,lag_intercept=True,knots='regional ERA5 daily temperature tertiles; fixed before fitting',sensitivity_configs=['lag14_df3x3','lag28_df3x3','lag21_df4x4']),humidity=dict(window='lag03',df=3,knots='regional ERA5 lag03 humidity tertiles; fixed before fitting'),calendar=dict(holiday='statutory_holiday',adjusted_workday=True,sensitivity='official_break'),matching_completeness='exclude entire event stratum if any required selected date/history is missing; report losses; never select replacement controls by exposure',variance=dict(primary='person-cluster robust sandwich',temporal_sensitivity='calendar-year-stratified whole-month block bootstrap; retain all events and control dates per sampled month; 200 draws, seed20261002; also review three-month blocks if serial dependence remains'),multiple_testing=dict(primary_family=15,method='Benjamini-Hochberg',report_raw_p=True,report_q=True,secondary='separate, labelled exploratory'),address_quality=dict(primary='verified reliable within about1km; GCJ02 converted to EPSG4326',sensitivity='unverified/fuzzy/coarse labelled separately; township-only references excluded from main'),subgroups=dict(age=['18-64','65+'],sex=['男','女'],season=dict(warm_months=[4,5,6,7,8,9],cold_months=[10,11,12,1,2,3]),formal_test='pollutant by subgroup interaction, no fixed subgroup main effect'),nonlinear_pollution=dict(df=3,knots='pooled eligible matched-date concentration tertiles fixed before coefficients are estimated',reference='pooled median',display='5th-95th percentile; do not infer sparse extremes'),two_pollutant=[['PM25','O3'],['PM10','O3']],pollution_distributed_lag=dict(max_lag=7,exposure='linear',lag_basis_df=3,status='exploratory; cumulative contrast distinct from moving-average coefficient'),components='not analyzed; available component product does not cover2023-2025',monitoring_flag_rule='daily onset count >max(20,3*annual median,median+6*1.4826*MAD); report count >max(100,annual median+6*1.4826*MAD),>=100 delayed>30d and fraction>=75%; descriptive delay list kept separately; flags only, no automatic exclusion',outcome_review='I20.0 primary angina coding candidate; other I20 not included in this acute primary category. Code/text conflicts require review; intervention itself does not establish acute onset.',dedup_task='explicitly deferred, no new identity grouping/merging/deletion performed by parallel tasks')
    target=O/'预设统计方案_v1.json';target.write_text(json.dumps(protocol,ensure_ascii=False,indent=2));receipt=dict(timestamp=protocol['prepared_at'],file=target.name,sha256=hashlib.sha256(target.read_bytes()).hexdigest(),status='核心参数本地预设，不是公开预注册；最终事件资格和去重尚未冻结',source_workbook_hashes=hashes.reset_index().to_dict('records'));(O/'方案版本与校验.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2))
    templates=pd.DataFrame([dict(outcome_id=k,outcome_name=n,pollutant=p,window='lag01',increment=10,OR=np.nan,lower=np.nan,upper=np.nan,p=np.nan,q=np.nan,n_events=np.nan,status='待最终事件、去重与定位核验后拟合') for k,n,_ in OUTCOMES for p in ['PM25','PM10','O3']]);save(templates,'15项主结果表模板.csv')
    import matplotlib;matplotlib.use('Agg');import matplotlib.pyplot as plt
    from matplotlib.font_manager import FontProperties
    font=FontProperties(fname='/System/Library/Fonts/STHeiti Medium.ttc');fig,axes=plt.subplots(5,1,figsize=(12,11),sharex=True)
    for ax,(key,name,prefs) in zip(axes,OUTCOMES):
        sub=monthly[(monthly.outcome_id==key)&(monthly.population=='既有候选事件（未新增去重）')];ax.plot(pd.to_datetime(sub.month),sub.n,marker='o',ms=3);ax.set_title(name,fontproperties=font,loc='left');ax.grid(alpha=.2);ax.axvline(pd.Timestamp('2023-04-01'),color='grey',ls='--',alpha=.6);ax.axvline(pd.Timestamp('2023-11-01'),color='grey',ls=':',alpha=.6)
    fig.suptitle('按预设ICD结局分组的逐月候选事件数',fontproperties=font);fig.text(.08,.012,'不稳定型心绞痛仅I20.0；未新增身份证去重；虚线标记2023年系统上线和联通月份。',fontproperties=font,fontsize=10);fig.tight_layout(rect=[0,.035,1,.97]);fig.savefig(O/'预设结局监测时间分布.png',dpi=160);plt.close(fig)
    result=dict(raw_records=len(raw),existing_candidate_events=len(events),code_defined_five_outcome_candidates=int((annotation.proposed_outcome!='').sum()),unstable_angina_candidates=int((annotation.proposed_outcome=='I20_0').sum()),code_text_conflicts=int((annotation.clinical_review_flag=='code_text_conflict').sum()),monitoring_flags=len(flagdf),new_identity_dedup_performed=False,real_association_models_fitted=0)
    (O/'并行任务汇总.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':main()
